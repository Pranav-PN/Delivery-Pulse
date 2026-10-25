# =============================================================================
# build_word_report.py – Compile DeliveryPulse Report to Microsoft Word (.docx)
# =============================================================================

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Style header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E78") # Deep Navy
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
            run.font.name = "Calibri"
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F2F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = "Calibri"
                
    if col_widths:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)
                
    doc.add_paragraph() # Spacer
    return table

def build_report():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles
    navy = RGBColor(31, 78, 120)
    dark_gray = RGBColor(50, 50, 50)
    
    # -------------------------------------------------------------------------
    # 1. TITLE PAGE
    # -------------------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(12)
    run_title = title_p.add_run("Predicting and Explaining Delivery Delays:\nA Probability and Statistical Learning Study of Food-Delivery Performance in RStudio")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = navy
    run_title.font.name = "Calibri"
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(36)
    run_sub = sub_p.add_run("DeliveryPulse: Operational Reliability, ETA-Breach Risk Modeling, and Dispatch Triage Optimization")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 100, 100)
    run_sub.font.name = "Calibri"
    
    # Metadata Box
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Course Code & Title:", "[Course Code Placeholder — e.g., BMAT202P Probability & Statistics]"),
        ("Student Investigator:", "[Student Name & Registration Number Placeholder]"),
        ("Faculty Guide / Mentor:", "[Faculty Name & Designation Placeholder]"),
        ("Academic Institution:", "[Department & University Name Placeholder]"),
        ("Date of Submission:", "October 2026")
    ]
    for idx, (label, val) in enumerate(meta_info):
        row = meta_table.rows[idx]
        row.cells[0].text = label
        row.cells[1].text = val
        set_cell_background(row.cells[0], "EBF1F5")
        set_cell_background(row.cells[1], "F8FAFC")
        set_cell_margins(row.cells[0], top=100, bottom=100, left=150, right=150)
        set_cell_margins(row.cells[1], top=100, bottom=100, left=150, right=150)
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(10)
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        
    doc.add_page_break()
    
    # -------------------------------------------------------------------------
    # Helper for headings
    # -------------------------------------------------------------------------
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = navy
        run.font.name = "Calibri"
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(46, 117, 182)
        run.font.name = "Calibri"
        return h

    def add_body(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.bold = True
            run_b.font.size = Pt(10.5)
            run_b.font.name = "Calibri"
        run = p.add_run(text)
        run.font.size = Pt(10.5)
        run.font.name = "Calibri"
        run.font.color.rgb = dark_gray
        return p

    def add_callout(text, bold_title="Key Finding:"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F2F6FA")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r_title = p.add_run(f"{bold_title} ")
        r_title.font.bold = True
        r_title.font.color.rgb = navy
        r_title.font.size = Pt(10)
        r_text = p.add_run(text)
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = dark_gray
        cell.width = Inches(6.5)
        doc.add_paragraph()

    def add_figure(img_path, caption):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(5.5))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            run_cap = p_cap.add_run(caption)
            run_cap.font.size = Pt(9)
            run_cap.font.italic = True
            run_cap.font.color.rgb = RGBColor(80, 80, 80)
            run_cap.font.name = "Calibri"

    # -------------------------------------------------------------------------
    # 2. INTRODUCTION
    # -------------------------------------------------------------------------
    add_h1("2. Introduction")
    add_h2("2.1 The Delivery-Delay Problem in Quick-Commerce Logistics")
    add_body("In modern on-demand food logistics and quick-commerce operations, meeting promised delivery windows is fundamental to platform viability. Last-mile fulfillment accounts for both the highest operational expenditure and the greatest volatility across the logistics supply chain. Environmental disruptions (fog, rain, storms), traffic bottlenecks, multi-order routing loads, and partner characteristics frequently lead to Estimated Time of Arrival (ETA) breaches. When food orders surpass the platform's standard 30-minute threshold, customer dissatisfaction surges, leading to app churn and financial losses from compensation credits. Treating delivery punctuality as a random Bernoulli trial allows platforms to predict ETA breaches proactively using probability theory and statistical learning.")
    
    add_h2("2.2 Dataset Description")
    add_body("This project utilizes the Food Delivery Dataset (Gaurav Malik, Kaggle), recording 45,593 delivery fulfillments across major metropolitan centers in India. Features include delivery personnel age and historical ratings, geographic coordinates (restaurant and delivery latitude/longitude), order timestamps, atmospheric conditions, traffic density, vehicle condition, multiple order loads, festival indicators, city types, and final delivery transit duration in minutes.")

    add_h2("2.3 Project Aim and Specific Objectives")
    add_body("To import the empirical food-delivery dataset, clean and engineer spatial transit metrics, analyze distributional properties and hypothesis tests, build logistic regression and Naive Bayes probabilistic classifiers to predict order delay risk, evaluate their predictive fidelity on held-out test data, and derive actionable triage recommendations for logistics dispatch control centers.", bold_prefix="Aim: ")
    
    add_body("Data Cleaning & Spatial Engineering: Parse raw records, convert demographic features, compute great-circle transit distances via Haversine's formula, filter erroneous GPS pings, and establish a binary delay threshold (Late: >30 min).", bold_prefix="• Objective 1 — ")
    add_body("Distribution Fitting & Confidence Intervals: Model delivery times with continuous Gamma and Lognormal distributions via Maximum Likelihood Estimation (MLE), compare AIC scores, and compute Clopper-Pearson exact and Wilson 95% confidence intervals.", bold_prefix="• Objective 2 — ")
    add_body("Hypothesis Testing Suite: Perform Welch two-sample t-tests and Chi-Square tests of independence with Benjamini-Hochberg False Discovery Rate corrections to identify statistically significant delay drivers.", bold_prefix="• Objective 3 — ")
    add_body("Probabilistic Machine Learning: Train multivariate Logistic Regression and Gaussian/Multinomial Naive Bayes classifiers on a 70% stratified training partition.", bold_prefix="• Objective 4 — ")
    add_body("Model Evaluation & Dispatch Triage: Evaluate classification metrics on the 30% held-out test set (Confusion Matrices, Sensitivity, Specificity, ROC/AUC) and evaluate a top-20% capacity triage intervention metric.", bold_prefix="• Objective 5 — ")

    # -------------------------------------------------------------------------
    # 3. STATISTICAL AND PROBABILITY FOUNDATIONS
    # -------------------------------------------------------------------------
    add_h1("3. Statistical and Probability Foundations")
    add_body("Delivery outcome is formulated as a discrete Bernoulli random variable Y ~ Bernoulli(p), where Y = 1 if Time taken > 30 minutes (Late) and Y = 0 if Time taken <= 30 minutes (On-Time). Across a batch of n independent orders, the number of delayed orders follows a Binomial distribution K ~ Binomial(n, p).", bold_prefix="3.1 Bernoulli & Binomial Process: ")
    add_body("To infer population parameter p from sample point estimate p_hat = k/n, we construct two 95% confidence intervals: the Clopper-Pearson exact interval (based on inverted cumulative binomial probabilities ensuring exact coverage) and the Wilson score interval (asymptotically optimal without over-conservatism).", bold_prefix="3.2 Confidence Intervals for Proportions: ")
    add_body("Delivery transit times T > 0 are continuous and right-skewed. We model T using Gamma(shape, rate) and Lognormal(meanlog, sdlog) distributions fitted via Maximum Likelihood Estimation (MLE). Goodness-of-fit is quantitatively adjudicated using the Akaike Information Criterion (AIC = 2k - 2*ln(L)).", bold_prefix="3.3 Parametric Distribution Modeling: ")
    add_body("Welch's two-sample t-test is employed to test for differences in continuous variables (distance, rider age, ratings) between On-Time and Late groups without assuming equal variances (heteroscedasticity-robust with Satterthwaite degrees of freedom).", bold_prefix="3.4 Welch's Two-Sample t-Test: ")
    add_body("Contingency tables evaluate independence between categorical predictors and delay outcomes using Pearson's Chi-Square statistic. Strength of association is quantified via Cramer's V, and multiple testing is controlled using the Benjamini-Hochberg (BH) False Discovery Rate procedure.", bold_prefix="3.5 Chi-Square & Cramer's V: ")
    add_body("Pearson's correlation measures linear dependency among continuous attributes. For distance vs. the binary Late indicator, this evaluates to the Point-Biserial correlation coefficient.", bold_prefix="3.6 Correlation Analysis: ")
    add_body("Logistic regression models the log-odds log(p/(1-p)) = beta_0 + sum(beta_j * X_j). Exponentiated coefficients exp(beta_j) represent odds ratios with 95% Wald confidence intervals. Naive Bayes calculates posterior class probabilities using Bayes' Theorem under conditional feature independence.", bold_prefix="3.7 Classification Models & ROC Analysis: ")

    # -------------------------------------------------------------------------
    # 4. ALGORITHMS USED
    # -------------------------------------------------------------------------
    add_h1("4. Algorithms Used")
    add_body("To prevent class distribution shift between partitions, stratified random sampling via caTools::sample.split partitions the 41,222 cleaned orders into 70% training (28,855 rows) and 30% testing (12,367 rows), ensuring exactly 30.62% late deliveries in both subsets.", bold_prefix="4.1 Stratified Train/Test Partitioning: ")
    add_body("Fitting the Generalized Linear Model (GLM) requires maximizing the binomial log-likelihood. R's glm() solves this via Iteratively Reweighted Least Squares (IRLS / Newton-Raphson), updating coefficient estimates until convergence.", bold_prefix="4.2 Iteratively Reweighted Least Squares (IRLS): ")
    add_body("e1071::naiveBayes computes class priors and feature likelihoods (class-conditional Gaussian parameters for continuous variables, frequency tables for factors), applying Bayes' rule to produce posterior probabilities.", bold_prefix="4.3 Naive Bayes Learning Algorithm: ")

    # -------------------------------------------------------------------------
    # 5. FUNCTIONS USED TABLE
    # -------------------------------------------------------------------------
    add_h1("5. Functions Used Across Scripts")
    add_body("The complete data pipeline, statistical tests, and machine learning models across scripts 01 through 06 utilize the following key R functions:")
    
    fn_headers = ["Function", "Package", "Statistical Purpose & Operational Role"]
    fn_data = [
        ["read.csv() / write.csv()", "utils (base)", "Loads raw/cleaned datasets and serializes processed tables."],
        ["trimws()", "base", "Strips whitespace to prevent phantom factor levels."],
        ["distHaversine()", "geosphere", "Computes spherical surface distance (km) between restaurant & delivery GPS."],
        ["complete.cases()", "stats (base)", "Filters complete observations across all modeling attributes."],
        ["binom.test()", "stats (base)", "Computes Clopper-Pearson exact 95% confidence interval for delay rate."],
        ["prop.test()", "stats (base)", "Computes Wilson score 95% confidence interval without continuity correction."],
        ["fitdist()", "fitdistrplus", "Fits Gamma and Lognormal continuous distributions using MLE."],
        ["denscomp(), qqcomp()", "fitdistrplus", "Generates empirical vs. fitted density, Q-Q, CDF, and P-P diagnostic plots."],
        ["pbinom()", "stats (base)", "Evaluates cumulative binomial probability for operational batch risk."],
        ["t.test()", "stats (base)", "Executes Welch two-sample t-tests with unequal group variances."],
        ["chisq.test()", "stats (base)", "Conducts Pearson Chi-Square tests of independence on contingency tables."],
        ["p.adjust(method='BH')", "stats (base)", "Controls False Discovery Rate across multiple hypothesis tests."],
        ["cor(), cor.test()", "stats (base)", "Calculates Pearson and point-biserial correlations with significance tests."],
        ["sample.split()", "caTools", "Executes 70/30 stratified random train/test partitioning."],
        ["glm(family=binomial)", "stats (base)", "Fits logistic regression models via Iteratively Reweighted Least Squares."],
        ["confint.default()", "stats (base)", "Calculates Wald 95% confidence intervals for odds ratio estimates."],
        ["naiveBayes()", "e1071", "Trains Naive Bayes classifier with Gaussian & multinomial likelihoods."],
        ["predict()", "stats / e1071", "Generates posterior probabilities and class assignments for test orders."],
        ["roc(), auc(), ci.auc()", "pROC", "Computes empirical ROC trajectories, AUC, and DeLong confidence intervals."]
    ]
    create_styled_table(doc, fn_headers, fn_data, [1.8, 1.2, 3.5])

    # -------------------------------------------------------------------------
    # 6. DATA ANALYSIS & EMPIRICAL FINDINGS
    # -------------------------------------------------------------------------
    add_h1("6. Data Analysis & Empirical Results")
    
    add_h2("6.1 Data Cleaning & Sample Verification (Script 01)")
    add_body("The raw dataset contained 45,593 records across 20 attributes. Cleaning involved stripping text tags from Time_taken(min), coercing 'NaN' tokens to true NAs, and calculating Haversine distances. Removing 431 anomalous coordinates (distance = 0 or > 50 km) and 3,940 missing values resulted in a final clean dataset of 41,222 complete observations. Of these, 12,622 orders (30.62%) exceeded 30 minutes, establishing the baseline late rate.")
    
    add_h2("6.2 Exploratory Data Analysis & Distribution Fitting (Script 02)")
    add_body("Clopper-Pearson Exact 95% CI: [0.3017, 0.3107]; Wilson Score 95% CI: [0.3018, 0.3107].", bold_prefix="Confidence Intervals: ")
    add_body("Fitting continuous distributions to transit duration revealed that the Gamma distribution (shape = 7.940, rate = 0.299, AIC = 298,377.9) substantially outperformed the Lognormal distribution (AIC = 299,176.8) by 798.9 AIC units.", bold_prefix="Parametric Distribution Comparison: ")
    add_body("For a random dispatch batch of n = 50 deliveries, the binomial probability that at least 40% (>= 20 orders) are delayed is P(K >= 20) = 0.1011 (10.11%).", bold_prefix="Binomial Batch Risk: ")
    
    add_figure("../plots/01_time_taken_fit.png", "Figure 1: Goodness-of-fit diagnostic plots (Density, Q-Q, CDF, P-P) demonstrating superior Gamma fit over Lognormal.")
    add_figure("../plots/02_late_distribution.png", "Figure 2: Delivery punctuality class distribution (69.4% On-Time vs. 30.6% Late).")
    add_figure("../plots/03_distance_by_late.png", "Figure 3: Delivery transit distance (km) stratified by punctuality outcome.")
    add_figure("../plots/04_ratings_by_late.png", "Figure 4: Delivery personnel ratings stratified by punctuality outcome.")
    add_figure("../plots/05_late_by_traffic.png", "Figure 5: Late delivery rates escalating monotonically across road traffic density levels.")
    add_figure("../plots/06_late_by_weather.png", "Figure 6: Late delivery proportions across atmospheric weather conditions.")

    add_h2("6.3 Hypothesis Testing Suite (Script 03)")
    add_body("Welch two-sample t-tests confirmed significant differences across all continuous features between on-time and late cohorts:")
    
    t_headers = ["Continuous Feature", "Mean On-Time", "Mean Late", "Difference", "t-Statistic", "df", "p-value", "95% CI"]
    t_data = [
        ["distance_km", "8.73 km", "12.19 km", "+3.46 km", "61.71", "25,756.1", "< 0.0001", "[3.35, 3.57]"],
        ["Delivery_person_Age", "28.76 yrs", "31.54 yrs", "+2.79 yrs", "47.48", "25,498.6", "< 0.0001", "[2.67, 2.90]"],
        ["Delivery_person_Ratings", "4.73", "4.41", "-0.32", "-80.17", "14,713.4", "< 0.0001", "[-0.32, -0.31]"]
    ]
    create_styled_table(doc, t_headers, t_data, [1.4, 0.8, 0.8, 0.8, 0.7, 0.7, 0.6, 0.7])
    
    add_body("Chi-Square tests of independence with Benjamini-Hochberg FDR adjustments ranked categorical predictors by Cramer's V:")
    
    chi_headers = ["Categorical Feature", "Chi-Square", "df", "Cramer's V", "p-value (Raw)", "p-value (BH Adj)", "Significance"]
    chi_data = [
        ["Road_traffic_density", "6031.71", "3", "0.3825", "< 0.0001", "< 0.0001", "Significant"],
        ["multiple_deliveries", "5886.09", "3", "0.3779", "< 0.0001", "< 0.0001", "Significant"],
        ["Weatherconditions", "2438.82", "5", "0.2432", "< 0.0001", "< 0.0001", "Significant"],
        ["Festival", "1926.41", "1", "0.2162", "< 0.0001", "< 0.0001", "Significant"],
        ["Vehicle_condition", "1341.86", "2", "0.1804", "< 0.0001", "< 0.0001", "Significant"],
        ["City", "1075.67", "2", "0.1615", "< 0.0001", "< 0.0001", "Significant"],
        ["Type_of_vehicle", "458.18", "2", "0.1054", "< 0.0001", "< 0.0001", "Significant"],
        ["Type_of_order", "0.88", "3", "0.0046", "0.8300", "0.8300", "Not Significant"]
    ]
    create_styled_table(doc, chi_headers, chi_data, [1.6, 0.8, 0.4, 0.8, 0.8, 0.9, 0.8])

    add_h2("6.4 Correlation Analysis (Script 04)")
    add_body("Pearson correlation between distance and delivery duration is r = +0.3220 (p < 0.0001), while rider rating exhibits an inverse correlation of r = -0.3597 (p < 0.0001). The point-biserial correlation between distance and the binary Late indicator is r_pb = +0.2837 (p < 0.0001). Inter-predictor correlations remained below |r| <= 0.115, verifying minimal multicollinearity.")
    add_figure("../plots/07_correlation_heatmap.png", "Figure 7: Pearson correlation heatmap among continuous operational variables.")

    add_h2("6.5 Predictive Modeling & Exponentiated Odds Ratios (Script 05)")
    add_body("Multivariate logistic regression was fitted on the 70% training split (N = 28,855). Key odds ratios and Wald confidence intervals:")
    
    or_headers = ["Feature", "Coefficient (Beta)", "Std. Error", "z-value", "Odds Ratio", "95% Wald CI", "p-value"]
    or_data = [
        ["distance_km", "+0.1463", "0.0045", "32.44", "1.1575", "[1.147, 1.168]", "< 0.0001"],
        ["Delivery_person_Age", "+0.1244", "0.0038", "32.64", "1.1325", "[1.124, 1.141]", "< 0.0001"],
        ["Delivery_person_Ratings", "-5.2833", "0.0896", "-58.94", "0.0051", "[0.004, 0.006]", "< 0.0001"],
        ["Road_traffic_densityLow", "-3.6294", "0.0891", "-40.72", "0.0265", "[0.022, 0.032]", "< 0.0001"],
        ["Road_traffic_densityMedium", "-0.9224", "0.0739", "-12.49", "0.3976", "[0.344, 0.460]", "< 0.0001"],
        ["WeatherconditionsSunny", "-3.2260", "0.0888", "-36.34", "0.0397", "[0.033, 0.047]", "< 0.0001"],
        ["multiple_deliveries1", "+0.4285", "0.0467", "9.17", "1.5350", "[1.401, 1.682]", "< 0.0001"],
        ["Type_of_vehiclemotorcycle", "+1.0627", "0.0832", "12.78", "2.8943", "[2.459, 3.407]", "< 0.0001"],
        ["CityUrban", "-0.5197", "0.0535", "-9.71", "0.5947", "[0.535, 0.661]", "< 0.0001"]
    ]
    create_styled_table(doc, or_headers, or_data, [1.6, 0.9, 0.7, 0.6, 0.7, 0.9, 0.7])

    add_h2("6.6 Test Set Evaluation & Operational Triage (Script 06)")
    add_body("Performance on the held-out test partition (N = 12,367) demonstrated high classification fidelity across both algorithms:")
    
    comp_headers = ["Model", "Test Accuracy", "Sensitivity (Recall)", "Specificity", "ROC / AUC"]
    comp_data = [
        ["Logistic Regression", "88.42%", "75.94%", "93.93%", "0.9386"],
        ["Naive Bayes Classifier", "88.61%", "69.82%", "96.90%", "0.9397"]
    ]
    create_styled_table(doc, comp_headers, comp_data, [2.0, 1.1, 1.2, 1.1, 1.1])
    
    add_figure("../plots/08_roc_curves.png", "Figure 8: Overlaid Receiver Operating Characteristic curves on held-out test data (Logistic AUC: 0.9386, Naive Bayes AUC: 0.9397).")
    
    add_callout(
        "If a dispatch operations team can only intervene on the top 20% riskiest in-flight orders (2,474 orders out of 12,367), the logistic regression model correctly captures 2,342 out of 3,787 actually late deliveries—recovering 61.84% of all delayed orders with a 3.09x efficiency lift over random auditing.",
        bold_title="Operational Triage Lift:"
    )

    # -------------------------------------------------------------------------
    # 7. SUMMARY & RECOMMENDATIONS
    # -------------------------------------------------------------------------
    add_h1("7. Summary of Results & Platform Recommendations")
    add_body("1. Dynamic ETA Buffering: In adverse traffic ('Jam') and weather (storms/fog), platforms should automatically append 10-15 minute buffers to customer arrival estimates, aligning expectations and avoiding contractual delay penalties.")
    add_body("2. Multi-Delivery Order Capping: Multiple deliveries per trip significantly elevate delay odds (OR = 1.535 for one extra stop). Dispatch algorithms should prohibit multi-order stacking during high traffic hours or adverse weather.")
    add_body("3. Proactive Dispatch Triage Engine: Real-time scoring using the logistic regression model allows operations control teams to monitor in-flight orders, intercepting the top 20% riskiest deliveries to achieve over 60% delay mitigation.")
    add_body("4. Experience-Based Rider Allocation: Higher-rated riders achieve substantially lower delivery times (r = -0.36, OR = 0.0051). Long-distance (>12 km) and peak-hour orders should be routed to riders with ratings >= 4.7.")

    # -------------------------------------------------------------------------
    # 8. LIMITATIONS & FUTURE SCOPE
    # -------------------------------------------------------------------------
    add_h1("8. Limitations and Future Scope")
    add_body("Limitations include the use of great-circle Haversine distances rather than street-level turn-by-turn routing topologies, and the absence of kitchen prep durations. Future work will integrate OpenStreetMap routing engines, test non-linear tree ensembles (Random Forests, XGBoost), and explore live radar weather telemetry.")

    # -------------------------------------------------------------------------
    # 9. REFERENCES
    # -------------------------------------------------------------------------
    add_h1("9. References")
    refs = [
        "1. Malik, G., 'Food Delivery Dataset,' Kaggle, 2022. Available: kaggle.com/datasets/gauravmalik26/food-delivery-dataset.",
        "2. Hosmer, D. W., Lemeshow, S., and Sturdivant, R. X., Applied Logistic Regression, 3rd ed., Hoboken, NJ: John Wiley & Sons, 2013.",
        "3. James, G., Witten, D., Hastie, T., and Tibshirani, R., An Introduction to Statistical Learning: with Applications in R, 2nd ed., New York, NY: Springer, 2021.",
        "4. Agresti, A., Categorical Data Analysis, 3rd ed., Hoboken, NJ: John Wiley & Sons, 2013.",
        "5. Benjamini, Y. and Hochberg, Y., 'Controlling the false discovery rate: a practical and powerful approach to multiple testing,' Journal of the Royal Statistical Society: Series B, vol. 57, no. 1, pp. 289–300, 1995.",
        "6. Delignette-Muller, M. L. and Dutang, C., 'fitdistrplus: An R Package for Fitting Distributions,' Journal of Statistical Software, vol. 64, no. 4, pp. 1–34, 2015.",
        "7. Fawcett, T., 'An introduction to ROC analysis,' Pattern Recognition Letters, vol. 27, no. 8, pp. 861–874, 2006.",
        "8. Robin, X. et al., 'pROC: an open-source package for R and S+ to analyze and compare ROC curves,' BMC Bioinformatics, vol. 12, no. 1, p. 77, 2011.",
        "9. R Core Team, R: A Language and Environment for Statistical Computing, Vienna, Austria: R Foundation for Statistical Computing, 2026."
    ]
    for r in refs:
        add_body(r, space_after=4)

    output_path = "report/DeliveryPulse_Report.docx"
    doc.save(output_path)
    print(f"Report compiled successfully: {output_path}")

if __name__ == "__main__":
    build_report()
